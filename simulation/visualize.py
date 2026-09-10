import os

import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook


# ============================================================
# POLARSENSE AI
# OCEAN DATA VISUALIZATION
# ============================================================


FILE_PATH = "data/ocean_simulations.xlsx"

OUTPUT_FOLDER = "data/plots"


# ============================================================
# CHECK FILE
# ============================================================

if not os.path.exists(FILE_PATH):

    print("ERROR:")
    print(f"Could not find {FILE_PATH}")
    print()
    print("Run ocean.py first.")

    exit()


# ============================================================
# FIND LATEST RUN
# ============================================================

workbook = load_workbook(
    FILE_PATH,
    read_only=True
)

run_sheets = []

for sheet in workbook.sheetnames:

    if sheet.startswith("Run_"):

        run_sheets.append(sheet)


if not run_sheets:

    print("No simulation runs found.")

    exit()


# Sort Run_001, Run_002, Run_003...

run_sheets.sort()

latest_run = run_sheets[-1]


print()
print("==============================================")
print("       POLARSENSE AI - DATA VISUALIZER")
print("==============================================")
print()

print(f"Latest simulation: {latest_run}")


# ============================================================
# READ EXCEL DATA
# ============================================================

data = pd.read_excel(
    FILE_PATH,
    sheet_name=latest_run
)


print(
    f"Observations loaded: {len(data)}"
)


# ============================================================
# CREATE PLOTS FOLDER
# ============================================================

os.makedirs(
    OUTPUT_FOLDER,
    exist_ok=True
)


# ============================================================
# CONVERT TIMESTAMP
# ============================================================

data["timestamp"] = pd.to_datetime(
    data["timestamp"]
)


# ============================================================
# PLOT 1
# TEMPERATURE
# ============================================================

plt.figure(
    figsize=(10, 5)
)

plt.plot(
    data["timestamp"],
    data["temperature_c"]
)

plt.title(
    f"Polar Ocean Temperature - {latest_run}"
)

plt.xlabel("Time")

plt.ylabel(
    "Temperature (°C)"
)

plt.xticks(
    rotation=45
)

plt.tight_layout()

plt.savefig(
    f"{OUTPUT_FOLDER}/temperature.png"
)

plt.close()


# ============================================================
# PLOT 2
# WIND SPEED
# ============================================================

plt.figure(
    figsize=(10, 5)
)

plt.plot(
    data["timestamp"],
    data["wind_speed_ms"]
)

plt.title(
    f"Wind Speed - {latest_run}"
)

plt.xlabel("Time")

plt.ylabel(
    "Wind Speed (m/s)"
)

plt.xticks(
    rotation=45
)

plt.tight_layout()

plt.savefig(
    f"{OUTPUT_FOLDER}/wind_speed.png"
)

plt.close()


# ============================================================
# PLOT 3
# WAVE HEIGHT
# ============================================================

plt.figure(
    figsize=(10, 5)
)

plt.plot(
    data["timestamp"],
    data["wave_height_m"]
)

plt.title(
    f"Wave Height - {latest_run}"
)

plt.xlabel("Time")

plt.ylabel(
    "Wave Height (m)"
)

plt.xticks(
    rotation=45
)

plt.tight_layout()

plt.savefig(
    f"{OUTPUT_FOLDER}/wave_height.png"
)

plt.close()


# ============================================================
# PLOT 4
# SEA ICE
# ============================================================

plt.figure(
    figsize=(10, 5)
)

plt.plot(
    data["timestamp"],
    data["ice_concentration_percent"]
)

plt.title(
    f"Sea Ice Concentration - {latest_run}"
)

plt.xlabel("Time")

plt.ylabel(
    "Ice Concentration (%)"
)

plt.xticks(
    rotation=45
)

plt.tight_layout()

plt.savefig(
    f"{OUTPUT_FOLDER}/ice_concentration.png"
)

plt.close()


# ============================================================
# PLOT 5
# OCEAN CURRENT
# ============================================================

plt.figure(
    figsize=(10, 5)
)

plt.plot(
    data["timestamp"],
    data["current_speed_ms"]
)

plt.title(
    f"Ocean Current Speed - {latest_run}"
)

plt.xlabel("Time")

plt.ylabel(
    "Current Speed (m/s)"
)

plt.xticks(
    rotation=45
)

plt.tight_layout()

plt.savefig(
    f"{OUTPUT_FOLDER}/current_speed.png"
)

plt.close()


# ============================================================
# PLOT 6
# BUOY TRAJECTORY
# ============================================================

plt.figure(
    figsize=(8, 8)
)

plt.plot(
    data["longitude"],
    data["latitude"]
)

plt.scatter(
    data["longitude"].iloc[0],
    data["latitude"].iloc[0],
    label="Start"
)

plt.scatter(
    data["longitude"].iloc[-1],
    data["latitude"].iloc[-1],
    label="End"
)

plt.title(
    f"Buoy Trajectory - {latest_run}"
)

plt.xlabel(
    "Longitude"
)

plt.ylabel(
    "Latitude"
)

plt.legend()

plt.grid(True)

plt.tight_layout()

plt.savefig(
    f"{OUTPUT_FOLDER}/buoy_trajectory.png"
)

plt.close()


# ============================================================
# SUMMARY
# ============================================================

print()
print("==============================================")
print("           VISUALIZATION COMPLETE")
print("==============================================")

print()
print("Generated plots:")

print("  1. temperature.png")
print("  2. wind_speed.png")
print("  3. wave_height.png")
print("  4. ice_concentration.png")
print("  5. current_speed.png")
print("  6. buoy_trajectory.png")

print()
print(f"Saved inside: {OUTPUT_FOLDER}")

print("==============================================")